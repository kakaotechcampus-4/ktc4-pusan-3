"""급식관리지원센터 배포 엑셀 급식표 → `MealPlanJSON` — docs/meal-plan/meal-plan-pipeline-v1.md §1.

표 안이 전부 텍스트라 OCR 이 없다. 셀을 그대로 읽고 모양만 맞춘다. 알레르기 번호는 손대지 않는다 —
`raw` 에 그대로 남겨 규칙(app/rules/allergen.py)이 뽑게 한다.

실측한 시트 배치 (6종 시트 공통):

    1행  기관명 | (기관명 값) | 2026년 9월 식단표 (만1~2세 일반형 식단) | ... | ▶발행처 : ...
    2행  요일   | 월 | 화 | 수 | 목 | 금 | 토          ← 여기서 요일 열을 정한다
    ---- 주 단위 블록이 반복된다 ----
         날짜   |    | 1  | 2  | 3  | 4  | 5           ← 값이 없는 요일은 그 주에 없는 날
         오전간식 | 달걀채소죽:①  | 복숭아/치즈:⑪/②  | ...
         점심   | 보리밥 | ...                          ← A열이 빈 다음 줄들은 같은 끼니가 이어진다
                | 애호박국:⑤⑥ | ...
         오후간식 | ...
         열량(Kcal)/단백질(g) | ...                     ← 건너뛴다
         날짜   | 7 | 8 | ... | 24:추석연휴 | ...       ← "N:메모" 는 휴일 등 메모
    ※알레르기 유발식품 ...                              ← 여기서부터 각주. 읽기를 멈춘다

석식형 시트는 끼니가 "저 녁" 하나이고 라벨에 공백이 섞인다("날 짜"). 공백을 지우고 비교한다.
"""

import io
import re
from dataclasses import dataclass, field

import openpyxl

from app.providers.meal_plan.schema import (
    Meal,
    MealDay,
    MealPlanJSON,
    MealPlanSource,
    MealType,
    MenuItem,
    UnparsedCell,
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_TITLE = re.compile(r"(\d{4})\s*년\s*(\d{1,2})\s*월")
_DAY_HEADER = re.compile(r"^(\d{1,2})(?::(.*))?$")  # "24" · "24:추석연휴"
_WEEKDAYS = ("월", "화", "수", "목", "금", "토", "일")
_MEAL_LABELS: dict[str, MealType] = {
    "아침": MealType.BREAKFAST,
    "조식": MealType.BREAKFAST,
    "오전간식": MealType.SNACK_AM,
    "점심": MealType.LUNCH,
    "중식": MealType.LUNCH,
    "오후간식": MealType.SNACK_PM,
    "저녁": MealType.DINNER,
    "석식": MealType.DINNER,
}
_SKIP_LABELS = ("열량", "기관명")
_FOOTER_MARKS = ("※", "★", "알레르기유발")
_EXCEL_ERRORS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NULL!", "#NUM!")


class MealPlanReadError(ValueError):
    """급식표로 읽을 수 없는 파일. 메시지에 셀 내용은 싣지 않는다."""


def list_menu_sheets(data: bytes) -> list[str]:
    """급식표로 보이는 시트 이름. 위 5행 안에 "요일" 행이 있으면 급식표다 (안내문 시트는 없다)."""
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return [
            sheet.title
            for sheet in workbook.worksheets
            if _weekday_row(list(sheet.iter_rows(min_row=1, max_row=5, values_only=True)))
            is not None
        ]
    finally:
        workbook.close()


class XlsxMealPlanReader:
    """`MealPlanReader` 구현체. `sheet` 를 주지 않으면 급식표로 보이는 첫 시트를 읽는다."""

    source = MealPlanSource.XLSX

    def __init__(self, sheet: str | None = None) -> None:
        self._sheet = sheet

    def read(self, data: bytes, *, mime_type: str) -> MealPlanJSON:
        if mime_type != XLSX_MIME:
            raise MealPlanReadError(f"엑셀(.xlsx)만 읽는다: {mime_type}")
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if self._sheet is not None:
                if self._sheet not in workbook.sheetnames:
                    raise MealPlanReadError("시트가 없다")
                rows = list(workbook[self._sheet].iter_rows(values_only=True))
            else:
                rows = next(
                    (
                        candidate
                        for sheet in workbook.worksheets
                        if _weekday_row(candidate := list(sheet.iter_rows(values_only=True)))
                        is not None
                    ),
                    None,
                )
                if rows is None:
                    raise MealPlanReadError("급식표 시트(요일 행)가 없다")
        finally:
            workbook.close()
        return parse_rows(rows)


@dataclass
class _DayCells:
    day: int
    note: str | None = None
    meals: dict[MealType, list[MenuItem]] = field(default_factory=dict)


def parse_rows(rows: list[tuple]) -> MealPlanJSON:
    """시트 하나의 셀 값(행 튜플 목록)을 급식표로 바꾼다.

    엑셀 파일이 없어도 테스트할 수 있게 읽기와 분리했다.
    """
    weekday_row = _weekday_row(rows)
    if weekday_row is None:
        raise MealPlanReadError("급식표 시트(요일 행)가 없다")
    year_month = _year_month(rows[: weekday_row + 1])
    institution = _institution(rows[0])
    day_columns = [
        index
        for index, value in enumerate(rows[weekday_row])
        if index > 0 and _norm(value) in _WEEKDAYS
    ]

    days: dict[int, _DayCells] = {}
    unparsed: list[UnparsedCell] = []
    block: dict[int, int] = {}  # 열 번호 → 그 주의 일자. 날짜 행마다 새로 만든다
    meal: MealType | None = None

    for row in rows[weekday_row + 1 :]:
        label = _norm(row[0] if row else None)
        if any(mark in label for mark in _FOOTER_MARKS):
            break
        if label.startswith("날짜"):
            block, meal = _read_day_header(row, day_columns, days), None
            continue
        if label.startswith(_SKIP_LABELS):
            continue
        if label:
            meal = _MEAL_LABELS.get(label)  # 모르는 라벨이면 그 줄들은 건너뛴다
            if meal is None:
                continue
        if meal is None or not block:
            continue
        for column, day in block.items():
            text = _text(row[column] if column < len(row) else None)
            if not text:
                continue
            if text in _EXCEL_ERRORS:
                unparsed.append(UnparsedCell(day=day, meal_type=meal, raw=text, why="엑셀 오류값"))
                continue
            days[day].meals.setdefault(meal, []).extend(_split_items(text))

    return MealPlanJSON(
        source=MealPlanSource.XLSX,
        year_month=year_month,
        institution_name=institution,
        days=[
            MealDay(
                day=cells.day,
                meals=[Meal(meal_type=kind, items=items) for kind, items in cells.meals.items()],
                note=cells.note,
            )
            for cells in sorted(days.values(), key=lambda c: c.day)
        ],
        unparsed=unparsed,
    )


def _read_day_header(row: tuple, day_columns: list[int], days: dict[int, _DayCells]) -> dict:
    """날짜 행 → {열: 일자}. 숫자가 아닌 머리(생일식단 등 견본 열)는 그 주에서 뺀다."""
    block: dict[int, int] = {}
    for column in day_columns:
        text = _text(row[column] if column < len(row) else None)
        matched = _DAY_HEADER.match(text) if text else None
        if not matched:
            continue
        day = int(matched.group(1))
        note = (matched.group(2) or "").strip() or None
        block[column] = day
        days.setdefault(day, _DayCells(day=day, note=note))
    return block


def _split_items(text: str) -> list[MenuItem]:
    """셀 하나 → 메뉴 항목. `복숭아/치즈:⑪/②` 처럼 이름과 번호가 같은 개수로 갈리면 항목을 나눈다.

    `찐고구마/우유:②` 처럼 개수가 다르면 나누지 않는다 — 어느 번호가 어느 것인지 지어내지 않는다.
    `raw` 는 나눈 항목도 원문 조각(이름:번호)을 그대로 잇는다. 번호를 고치거나 지우지 않는다.
    """
    name_part, colon, code_part = text.partition(":")
    names = [part.strip() for part in name_part.split("/")]
    codes = [part.strip() for part in code_part.split("/")] if colon else []
    if len(names) > 1 and len(names) == len(codes):
        return [MenuItem(name=name, raw=f"{name}:{code}") for name, code in zip(names, codes)]
    return [MenuItem(name=name_part.strip(), raw=text)]


def _weekday_row(rows: list[tuple]) -> int | None:
    for index, row in enumerate(rows):
        if row and _norm(row[0]) == "요일":
            return index
    return None


def _year_month(rows: list[tuple]) -> str:
    for row in rows:
        for value in row:
            matched = _TITLE.search(str(value)) if value is not None else None
            if matched:
                return f"{int(matched.group(1)):04d}-{int(matched.group(2)):02d}"
    raise MealPlanReadError("제목에서 'YYYY년 M월' 을 찾지 못했다")


def _institution(first_row: tuple) -> str | None:
    """1행이 `기관명 | 값` 이면 그 값. 센터 배포본은 비어 있다."""
    if first_row and _norm(first_row[0]) == "기관명" and len(first_row) > 1:
        return _text(first_row[1]) or None
    return None


def _norm(value: object) -> str:
    return re.sub(r"\s+", "", str(value)) if value is not None else ""


def _text(value: object) -> str:
    """셀 값 → 문자열. 비었거나 0 이면 빈 문자열 (센터 배포본은 빈 칸을 0 으로 채운다)."""
    if value is None or value == 0 or value == "0":
        return ""
    return str(value).strip()
